import asyncio
import sqlite3
import tempfile
import unittest
from contextlib import closing
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from career_agent import main
from career_agent.models import ModelCatalog
from career_agent.provider import Provider
from career_agent.runtime import AgentRuntime
from career_agent.store import Store


class InterviewApiTests(unittest.TestCase):
    def setUp(self):
        temp_root = Path(__file__).parent / ".tmp"
        temp_root.mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=temp_root)
        self.original_store, self.original_runtime = main.store, main.runtime
        main.store = Store(Path(self.temp.name) / "api.sqlite3")
        provider = Provider()
        provider.mode = "demo"
        main.runtime = AgentRuntime(main.store, provider)
        self.client = TestClient(main.app)

    def tearDown(self):
        self.client.close()
        main.store, main.runtime = self.original_store, self.original_runtime
        self.temp.cleanup()

    def test_config_shows_missing_model_without_exposing_key(self):
        with patch.object(main.provider, "mode", "openai_compatible"), patch.object(main.provider, "model", ""):
            config = self.client.get("/api/config").json()
        self.assertEqual(config["config_error"], "请设置 MODEL_NAME")
        self.assertNotIn("api_key", config)

    def test_api_responses_are_not_cached_or_sniffed(self):
        response = self.client.get("/api/config")
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertEqual(response.headers["x-content-type-options"], "nosniff")
        self.assertEqual(response.headers["referrer-policy"], "no-referrer")
        page = self.client.get("/")
        self.assertNotIn("no-store", page.headers.get("cache-control", ""))

    def test_web_ui_exposes_memory_candidate_review_and_read_only_tool_state(self):
        page = self.client.get("/").text
        script = self.client.get("/static/app.js").text
        self.assertIn('id="settings-memory-candidates"', page)
        self.assertIn('id="settings-memory-candidate-list"', page)
        self.assertIn('id="settings-tool-list"', page)
        self.assertIn("非只读工具当前不可授权", page)
        self.assertIn("/memory-candidates?session_id=", script)
        self.assertIn("/confirm`, { method: \"POST\" }", script)
        self.assertIn("checkbox.disabled = !readOnly", script)
        self.assertIn("candidate-content-", script)

    def test_session_selects_configured_model_and_rejects_unknown(self):
        path = Path(self.temp.name) / "models.json"
        path.write_text('[{"id":"local","provider":"openai_compatible","model":"test-model",'
                        '"base_url":"http://127.0.0.1:1234/v1","supports_tools":true}]', encoding="utf-8")
        with patch.object(main, "model_catalog", ModelCatalog(path)):
            models = self.client.get("/api/models").json()
            self.assertEqual(models[1]["id"], "local")
            chat = self.client.post("/api/chat/sessions", json={"model_id": "local"})
            self.assertEqual(chat.status_code, 200)
            self.assertEqual(chat.json()["model_id"], "local")
            claimed = self.client.post(f"/api/chat/sessions/{chat.json()['id']}/messages/claim", json={
                "text": "固定模型",
            }).json()
            self.assertEqual(claimed["model_id"], "local")
            interview = self.client.post("/api/interview/sessions", json={
                "goal": "后端开发", "model_id": "local",
            })
            self.assertEqual(interview.json()["model_id"], "local")
            self.assertEqual(self.client.post("/api/chat/sessions", json={
                "model_id": "missing",
            }).status_code, 422)

    def test_long_term_memory_search_and_delete(self):
        created = self.client.post("/api/memories", json={"content": "熟悉 Python 和 SQL"})
        self.assertEqual(created.status_code, 200)
        memory_id = created.json()["id"]
        self.assertEqual(len(self.client.get("/api/memories", params={"q": "Python"}).json()), 1)
        self.assertEqual(self.client.get("/api/memories", params={"q": "%"}).json(), [])
        self.assertEqual(self.client.post("/api/memories", json={"content": "  "}).status_code, 422)
        self.assertEqual(self.client.delete(f"/api/memories/{memory_id}").status_code, 200)
        self.assertEqual(self.client.get("/api/memories").json(), [])
        self.assertEqual(self.client.delete(f"/api/memories/{memory_id}").status_code, 404)

    def test_long_term_memory_ignores_repeats_and_respects_the_cap(self):
        first = self.client.post("/api/memories", json={"content": "熟悉  Python   和 SQL"}).json()
        self.assertIs(first["created"], True)
        self.assertEqual(first["content"], "熟悉 Python 和 SQL")
        repeat = self.client.post("/api/memories", json={"content": "熟悉 Python 和 SQL"})
        self.assertEqual(repeat.json()["id"], first["id"])
        self.assertIs(repeat.json()["created"], False)
        self.assertEqual(len(self.client.get("/api/memories").json()), 1)

        with patch("career_agent.store.MAX_MEMORIES", 1):
            over = self.client.post("/api/memories", json={"content": "另一条记忆"})
        self.assertEqual(over.status_code, 422)
        self.assertIn("上限", over.json()["detail"])
        self.assertEqual(len(self.client.get("/api/memories").json()), 1)

    def test_memory_candidate_review_and_confirmation(self):
        chat = self.client.post("/api/chat/sessions").json()
        message = self.client.post(f"/api/chat/sessions/{chat['id']}/messages", json={
            "text": "请记住我偏好后端开发岗位",
        })
        self.assertEqual(message.status_code, 200, message.text)
        source = self.client.get(f"/api/chat/sessions/{chat['id']}").json()["messages"][0]

        created = self.client.post(f"/api/chat/sessions/{chat['id']}/memory-candidates", json={
            "content": "偏好后端开发岗位",
            "source_message_ids": [source["id"]],
        })
        self.assertEqual(created.status_code, 200, created.text)
        candidate = created.json()
        self.assertEqual(candidate["status"], "pending")
        self.assertEqual(candidate["sources"][0]["id"], source["id"])
        self.assertEqual(self.client.get("/api/memories", params={
            "q": "后端开发",
        }).json(), [])
        pending = self.client.get("/api/memory-candidates", params={
            "session_id": chat["id"], "status": "pending",
        }).json()
        self.assertEqual([item["id"] for item in pending], [candidate["id"]])

        edited = self.client.put(f"/api/memory-candidates/{candidate['id']}", json={
            "content": "长期偏好：后端开发岗位",
        }).json()
        self.assertEqual(edited["content"], "长期偏好：后端开发岗位")
        confirmed = self.client.post(f"/api/memory-candidates/{candidate['id']}/confirm")
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertEqual(confirmed.json()["status"], "confirmed")
        self.assertEqual(self.client.get("/api/memories", params={
            "q": "后端开发",
        }).json()[0]["content"], "长期偏好：后端开发岗位")
        self.assertEqual(self.client.delete(f"/api/memory-candidates/{candidate['id']}").status_code, 409)

        memory_id = confirmed.json()["memory"]["id"]
        self.client.delete(f"/api/memories/{memory_id}")
        self.assertIsNone(self.client.get(f"/api/memory-candidates/{candidate['id']}").json()["memory_id"])

        removable = self.client.post(f"/api/chat/sessions/{chat['id']}/memory-candidates", json={
            "content": "稍后删除的候选",
            "source_message_ids": [source["id"]],
        }).json()
        self.assertEqual(self.client.delete(
            f"/api/memory-candidates/{removable['id']}"
        ).status_code, 200)
        self.assertEqual(self.client.get(
            f"/api/memory-candidates/{removable['id']}"
        ).status_code, 404)

        other = self.client.post("/api/chat/sessions").json()
        rejected = self.client.post(f"/api/chat/sessions/{other['id']}/memory-candidates", json={
            "content": "错误来源",
            "source_message_ids": [source["id"]],
        })
        self.assertEqual(rejected.status_code, 422)

    def test_session_titles_are_display_only_and_can_be_reset(self):
        chat = self.client.post("/api/chat/sessions").json()
        pending = self.client.post(f"/api/chat/sessions/{chat['id']}/messages/claim", json={
            "text": "讨论后端岗位",
        }).json()
        saved = self.client.put(f"/api/chat/sessions/{chat['id']}/title", json={
            "title": "  后端求职准备  ",
        })
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json()["title"], "后端求职准备")
        self.assertEqual(saved.json()["version"], pending["version"])
        self.assertEqual(self.client.get("/api/chat/sessions").json()[0]["title"], "后端求职准备")
        self.assertNotIn("后端求职准备", str(main.runtime.build_context(saved.json())))
        self.assertIn("会话名称：后端求职准备", self.client.get(
            f"/api/chat/sessions/{chat['id']}/export.md"
        ).text)
        self.assertEqual(self.client.post(f"/api/chat/sessions/{chat['id']}/messages/run", json={
            "version": pending["version"],
        }).status_code, 200)
        self.assertEqual(self.client.put(f"/api/interview/sessions/{chat['id']}/title", json={
            "title": "错误入口",
        }).status_code, 404)
        self.assertEqual(self.client.put(f"/api/chat/sessions/{chat['id']}/title", json={
            "title": "过长" * 41,
        }).status_code, 422)
        reset = self.client.put(f"/api/chat/sessions/{chat['id']}/title", json={"title": "  "})
        self.assertEqual(reset.json()["title"], "")
        one_line = self.client.put(f"/api/chat/sessions/{chat['id']}/title", json={
            "title": "第一行\n第二行",
        })
        self.assertEqual(one_line.json()["title"], "第一行 第二行")

        interview = self.client.post("/api/interview/sessions", json={"goal": "算法岗位"}).json()
        renamed = self.client.put(f"/api/interview/sessions/{interview['id']}/title", json={
            "title": "第一次模拟",
        }).json()
        self.assertEqual(renamed["title"], "第一次模拟")
        self.assertEqual(renamed["goal"], "算法岗位")
        self.assertEqual(self.client.get("/api/interview/sessions").json()[0]["title"], "第一次模拟")

    def test_session_search_finds_messages_and_interview_feedback(self):
        chat = self.client.post("/api/chat/sessions").json()
        chat_id = chat["id"]
        self.client.put(f"/api/chat/sessions/{chat_id}/memory", json={"text": "PRIVATE_CHAT_NOTE"})
        self.client.post(f"/api/chat/sessions/{chat_id}/messages/claim", json={
            "text": "后端项目中的索引优化",
        })
        other = self.client.post("/api/chat/sessions").json()
        self.client.put(f"/api/chat/sessions/{other['id']}/title", json={"title": "100XX 岗位"})
        self.client.put(f"/api/chat/sessions/{chat_id}/title", json={"title": "100%_ 岗位"})
        for query in ("索引优化", "100%_", chat_id[:8]):
            matches = self.client.get("/api/chat/sessions", params={"q": query}).json()
            self.assertEqual([item["id"] for item in matches], [chat_id])
            self.assertNotIn("索引优化", str(matches))
            self.assertNotIn("PRIVATE_CHAT_NOTE", str(matches))
        self.assertEqual(self.client.get("/api/chat/sessions", params={"q": "不存在的内容"}).json(), [])
        self.assertEqual(self.client.get("/api/chat/sessions", params={"q": "甲" * 101}).status_code, 422)

        interview = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "max_questions": 1,
        }).json()
        interview_id = interview["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
            "text": "我优化了数据库连接池",
        })
        with main.store.connect() as db:
            db.execute("UPDATE interview_items SET feedback = ? WHERE session_id = ?", ("反馈专属关键词", interview_id))
        for query in ("连接池", "反馈专属关键词", "后端开发"):
            matches = self.client.get("/api/interview/sessions", params={"q": query}).json()
            self.assertEqual([item["id"] for item in matches], [interview_id])
        self.assertEqual(self.client.get("/api/chat/sessions", params={"q": "反馈专属关键词"}).json(), [])

    def test_chat_memory_is_editable_and_invalidates_old_run(self):
        chat_id = self.client.post("/api/chat/sessions").json()["id"]
        pending = self.client.post(f"/api/chat/sessions/{chat_id}/messages/claim", json={
            "text": "旧问题",
        }).json()
        saved = self.client.put(f"/api/chat/sessions/{chat_id}/memory", json={
            "text": "  优先讨论后端岗位  ",
        })
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(saved.json()["text"], "优先讨论后端岗位")
        view = self.client.get(f"/api/chat/sessions/{chat_id}").json()
        self.assertEqual(view["session"]["chat_memory"], "优先讨论后端岗位")
        self.assertEqual(view["session"]["status"], "active")
        self.assertTrue(self.client.post(f"/api/chat/sessions/{chat_id}/messages/run", json={
            "version": pending["version"],
        }).json()["interrupted"])
        self.assertIn("优先讨论后端岗位", self.client.get(
            f"/api/chat/sessions/{chat_id}/export.md"
        ).text)
        self.assertEqual(self.client.put(f"/api/chat/sessions/{chat_id}/memory", json={
            "text": "甲" * 2001,
        }).status_code, 422)
        self.client.put(f"/api/chat/sessions/{chat_id}/memory", json={"text": ""})
        self.assertEqual(self.client.get(f"/api/chat/sessions/{chat_id}").json()["session"]["chat_memory"], "")

    def test_chat_retry_reuses_unanswered_message(self):
        chat_id = self.client.post("/api/chat/sessions").json()["id"]
        self.assertEqual(self.client.post(f"/api/chat/sessions/{chat_id}/retry").status_code, 409)
        material = self.client.post("/api/materials", json={
            "kind": "project", "title": "项目", "content": "缓存命中率提升",
        }).json()
        text = f"请解释材料 {material['id']}"
        pending = self.client.post(f"/api/chat/sessions/{chat_id}/messages/claim", json={
            "text": text,
        }).json()
        main.store.fail_run(chat_id, pending["version"], "chat", "model_unavailable")
        result = self.client.post(f"/api/chat/sessions/{chat_id}/retry")
        self.assertEqual(result.status_code, 200, result.text)
        self.assertIn("缓存命中率提升", result.json()["message"])
        messages = self.client.get(f"/api/chat/sessions/{chat_id}").json()["messages"]
        self.assertEqual([message["role"] for message in messages], ["user", "assistant"])
        self.assertEqual(messages[0]["content"], text)
        self.assertEqual(self.client.post(f"/api/chat/sessions/{chat_id}/retry").status_code, 409)

    def test_chat_followup_claim_cancels_inflight_run(self):
        class WaitingProvider:
            def __init__(self):
                self.started = asyncio.Event()
                self.cancelled = False

            async def complete(self, messages, tools):
                self.started.set()
                try:
                    await asyncio.Event().wait()
                except asyncio.CancelledError:
                    self.cancelled = True
                    raise

        async def scenario():
            provider = WaitingProvider()
            chat = main.store.create("chat")
            first = main.store.claim_run(chat["id"], "chat", "active", "旧问题")
            task = asyncio.create_task(AgentRuntime(main.store, provider).run(first))
            await provider.started.wait()
            second = await asyncio.to_thread(
                main.claim_chat_message, chat["id"], main.TextInput(text="补充问题")
            )
            with self.assertRaises(asyncio.CancelledError):
                await task
            return provider, first, second

        provider, first, second = asyncio.run(scenario())
        self.assertTrue(provider.cancelled)
        self.assertGreater(second["version"], first["version"])
        self.assertEqual(main.store.get(first["id"], "chat")["status"], "running")
        self.assertEqual([item["content"] for item in main.store.messages(first["id"], "chat")],
                         ["旧问题", "补充问题"])

    def test_chat_interrupt_cancels_inflight_run(self):
        class WaitingProvider:
            def __init__(self):
                self.started = asyncio.Event()
                self.cancelled = False

            async def complete(self, messages, tools):
                self.started.set()
                try:
                    await asyncio.Event().wait()
                except asyncio.CancelledError:
                    self.cancelled = True
                    raise

        async def scenario():
            provider = WaitingProvider()
            chat = main.store.create("chat")
            pending = main.store.claim_run(chat["id"], "chat", "active", "问题")
            task = asyncio.create_task(AgentRuntime(main.store, provider).run(pending))
            await provider.started.wait()
            interrupted = await asyncio.to_thread(main.interrupt_chat, chat["id"])
            with self.assertRaises(asyncio.CancelledError):
                await task
            return provider, interrupted

        provider, interrupted = asyncio.run(scenario())
        self.assertTrue(provider.cancelled)
        self.assertEqual(interrupted["status"], "active")
        self.assertEqual(main.store.get(interrupted["id"], "chat")["status"], "active")

    def test_backup_contains_committed_business_data(self):
        material = self.client.post("/api/materials", json={
            "kind": "resume", "title": "我的简历", "content": "备份中的经历",
        }).json()
        chat = self.client.post("/api/chat/sessions").json()
        self.client.post(f"/api/chat/sessions/{chat['id']}/messages", json={"text": "保存这条消息"})
        backup = self.client.get("/api/data/backup")
        self.assertEqual(backup.status_code, 200)
        self.assertIn("attachment", backup.headers["content-disposition"])
        backup_path = Path(self.temp.name) / "export.sqlite3"
        backup_path.write_bytes(backup.content)
        with closing(sqlite3.connect(backup_path)) as db:
            self.assertEqual(db.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertEqual(db.execute("SELECT content FROM materials WHERE id = ?", (material["id"],)).fetchone()[0], "备份中的经历")
            self.assertEqual(db.execute("SELECT content FROM messages WHERE session_id = ? AND role = 'user'", (chat["id"],)).fetchone()[0], "保存这条消息")

    def test_material_to_report_to_practice(self):
        material = self.client.post("/api/materials", json={
            "kind": "resume", "title": "项目简历", "content": "我负责性能优化。",
        }).json()
        claim = self.client.post("/api/claims", json={
            "text": "接口延迟下降 40%", "material_id": material["id"],
        }).json()
        denied = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "claim_ids": [claim["id"]],
        })
        self.assertEqual(denied.status_code, 422)
        self.assertEqual(self.client.post(f"/api/claims/{claim['id']}/confirm").status_code, 200)
        interview = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "material_ids": [material["id"]], "claim_ids": [claim["id"]],
        }).json()
        session_id = interview["id"]
        self.assertEqual(self.client.post(f"/api/interview/sessions/{session_id}/start").status_code, 200)
        for number in range(3):
            reply = self.client.post(
                f"/api/interview/sessions/{session_id}/answers", json={"text": f"回答 {number}"}
            )
            self.assertEqual(reply.status_code, 200)
        report = self.client.get(f"/api/interview/sessions/{session_id}/report").json()
        self.assertEqual(report["status"], "finished")
        self.assertEqual(len(report["items"]), 3)
        self.assertEqual(report["items"][0]["answer"], "回答 0")
        self.assertEqual(report["items"][0]["evidence"], ["回答 0"])
        practice = self.client.post(
            f"/api/interview/sessions/{session_id}/practice", json={"question_number": 2}
        ).json()
        practice_id = practice["id"]
        self.assertEqual(self.client.post(f"/api/interview/sessions/{practice_id}/start").status_code, 200)
        self.assertEqual(self.client.post(
            f"/api/interview/sessions/{practice_id}/answers", json={"text": "复练回答"}
        ).status_code, 200)
        practice_report = self.client.get(f"/api/interview/sessions/{practice_id}/report").json()
        self.assertEqual(practice_report["status"], "finished")
        self.assertEqual(len(practice_report["items"]), 1)
        self.assertEqual(practice_report["practice_source"]["question_number"], 2)
        self.assertEqual(practice_report["practice_source"]["answer"], "回答 1")
        self.assertEqual(len(self.client.get(
            f"/api/interview/sessions/{practice_id}"
        ).json()["resources"]["materials"]), 1)

    def test_chat_uses_confirmed_claims_and_explicit_resources(self):
        claim = self.client.post("/api/claims", json={"text": "我优化了查询性能"}).json()
        material = self.client.post("/api/materials", json={
            "kind": "project", "title": "查询项目", "content": "使用索引减少了扫描行数。",
        }).json()
        chat = self.client.post("/api/chat/sessions").json()
        chat_id = chat["id"]
        first_response = self.client.post(f"/api/chat/sessions/{chat_id}/messages", json={
            "text": "请整理我的已确认经历",
        })
        self.assertEqual(first_response.status_code, 200, first_response.text)
        first = first_response.json()
        self.assertIn("没有已确认", first["message"])
        self.client.post(f"/api/claims/{claim['id']}/confirm")
        confirmed = self.client.post(f"/api/chat/sessions/{chat_id}/messages", json={
            "text": "请整理我的已确认经历",
        }).json()
        self.assertIn("我优化了查询性能", confirmed["message"])
        named_material = self.client.post(f"/api/chat/sessions/{chat_id}/messages", json={
            "text": f"请解释材料 {material['id']}",
        }).json()
        self.assertIn("使用索引减少了扫描行数", named_material["message"])
        interview = self.client.post("/api/interview/sessions", json={"goal": "后端开发"}).json()
        interview_id = interview["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        for _ in range(3):
            self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
                "text": "PRIVATE_INTERVIEW_ANSWER",
            })
        discussed = self.client.post(f"/api/chat/sessions/{chat_id}/messages", json={
            "text": f"请解释面试 {interview_id} 的反馈",
        }).json()
        self.assertIn("面试总结", discussed["message"])
        self.assertNotIn("PRIVATE_INTERVIEW_ANSWER", discussed["message"])
        interview_view = self.client.get(f"/api/interview/sessions/{interview_id}").json()
        self.assertNotIn("请整理我的已确认经历", str(interview_view["messages"]))

    def test_edit_and_revoke_preserve_old_interview_resources(self):
        material = self.client.post("/api/materials", json={
            "kind": "project", "title": "旧项目", "content": "旧版实施细节",
        }).json()
        claim = self.client.post("/api/claims", json={
            "text": "旧版个人贡献", "material_id": material["id"],
        }).json()
        self.client.post(f"/api/claims/{claim['id']}/confirm")
        interview = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "material_ids": [material["id"]], "claim_ids": [claim["id"]],
        }).json()
        revised = self.client.put(f"/api/materials/{material['id']}", json={
            "title": "新版项目", "content": "新版实施细节",
        }).json()
        self.assertNotEqual(revised["id"], material["id"])
        self.assertEqual([item["id"] for item in self.client.get("/api/materials").json()], [revised["id"]])
        self.assertEqual(self.client.get("/api/claims").json(), [])
        self.assertEqual(self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "material_ids": [material["id"]],
        }).status_code, 422)
        self.assertEqual(self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "claim_ids": [claim["id"]],
        }).status_code, 422)
        old_view = self.client.get(f"/api/interview/sessions/{interview['id']}").json()
        self.assertEqual(old_view["resources"]["materials"][0]["title"], "旧项目")
        self.assertEqual(old_view["resources"]["claims"][0]["text"], "旧版个人贡献")
        self.assertEqual(self.client.delete(f"/api/materials/{revised['id']}").status_code, 200)
        self.assertEqual(self.client.get("/api/materials").json(), [])

    def test_edited_claim_requires_new_confirmation(self):
        claim = self.client.post("/api/claims", json={"text": "原贡献"}).json()
        self.client.post(f"/api/claims/{claim['id']}/confirm")
        edited = self.client.put(f"/api/claims/{claim['id']}", json={"text": "修正贡献"}).json()
        self.assertEqual(edited["status"], "draft")
        self.assertEqual(self.client.post("/api/interview/sessions", json={
            "goal": "开发", "claim_ids": [edited["id"]],
        }).status_code, 422)
        self.assertEqual(self.client.post(f"/api/claims/{edited['id']}/confirm").status_code, 200)
        self.assertEqual(self.client.delete(f"/api/claims/{edited['id']}").status_code, 200)
        self.assertEqual(self.client.get("/api/claims").json(), [])

    def test_chat_followup_supersedes_older_run_without_losing_messages(self):
        chat_id = self.client.post("/api/chat/sessions").json()["id"]
        first = self.client.post(f"/api/chat/sessions/{chat_id}/messages/claim", json={
            "text": "先讨论岗位",
        }).json()
        second = self.client.post(f"/api/chat/sessions/{chat_id}/messages/claim", json={
            "text": "补充：重点是后端开发",
        }).json()
        old = self.client.post(f"/api/chat/sessions/{chat_id}/messages/run", json={
            "version": first["version"],
        }).json()
        self.assertTrue(old["interrupted"])
        latest = self.client.post(f"/api/chat/sessions/{chat_id}/messages/run", json={
            "version": second["version"],
        }).json()
        self.assertIn("补充：重点是后端开发", latest["message"])
        messages = self.client.get(f"/api/chat/sessions/{chat_id}").json()["messages"]
        self.assertEqual([item["role"] for item in messages], ["user", "user", "assistant"])
        self.assertEqual([item["content"] for item in messages[:2]], ["先讨论岗位", "补充：重点是后端开发"])

    def test_interview_answer_correction_keeps_history_and_latest_feedback(self):
        interview_id = self.client.post("/api/interview/sessions", json={"goal": "开发"}).json()["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        first = self.client.post(f"/api/interview/sessions/{interview_id}/answers/claim", json={
            "text": "最初回答", "question_number": 1,
        }).json()
        corrected = self.client.post(f"/api/interview/sessions/{interview_id}/answers/claim", json={
            "text": "更正后的回答", "question_number": 1,
        }).json()
        old = self.client.post(f"/api/interview/sessions/{interview_id}/answers/run", json={
            "version": first["version"],
        }).json()
        self.assertTrue(old["interrupted"])
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/answers/run", json={
            "version": corrected["version"],
        }).status_code, 200)
        stale = self.client.post(f"/api/interview/sessions/{interview_id}/answers/claim", json={
            "text": "迟到的第 1 题答案", "question_number": 1,
        })
        self.assertEqual(stale.status_code, 409)
        report = self.client.get(f"/api/interview/sessions/{interview_id}/report").json()
        self.assertEqual(report["items"][0]["answer"], "更正后的回答")
        self.assertEqual(report["items"][0]["evidence"], ["更正后的回答"])
        messages = self.client.get(f"/api/interview/sessions/{interview_id}").json()["messages"]
        self.assertIn("最初回答", [item["content"] for item in messages])
        self.assertIn("更正后的回答", [item["content"] for item in messages])
        exported = self.client.get(f"/api/interview/sessions/{interview_id}/export.md").text
        self.assertIn("最初回答", exported)
        self.assertIn("更正后的回答", exported)

    def test_interview_respects_selected_question_count(self):
        self.assertEqual(self.client.post("/api/interview/sessions", json={
            "goal": "开发", "max_questions": 0,
        }).status_code, 422)
        interview = self.client.post("/api/interview/sessions", json={
            "goal": "开发", "max_questions": 2,
        }).json()
        interview_id = interview["id"]
        self.assertEqual(interview["max_questions"], 2)
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={"text": "第一题回答"})
        final = self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
            "text": "第二题回答",
        }).json()
        self.assertEqual(final["session"]["status"], "finished")
        self.assertEqual(len(self.client.get(f"/api/interview/sessions/{interview_id}/report").json()["items"]), 2)
        self.assertTrue(self.client.get("/api/interview/sessions").json()[0]["has_feedback"])
        practice = self.client.post(f"/api/interview/sessions/{interview_id}/practice", json={
            "question_number": 1,
        }).json()
        self.assertEqual(practice["max_questions"], 1)

    def test_interview_skip_records_unanswered_question_and_continues(self):
        interview_id = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "max_questions": 2,
        }).json()["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/skip", json={
            "question_number": 2,
        }).status_code, 409)
        skipped = self.client.post(f"/api/interview/sessions/{interview_id}/skip", json={
            "question_number": 1,
        })
        self.assertEqual(skipped.status_code, 200, skipped.text)
        self.assertEqual(skipped.json()["session"]["question_count"], 2)
        self.assertEqual(skipped.json()["session"]["status"], "awaiting_answer")
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/skip", json={
            "question_number": 1,
        }).status_code, 409)
        self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
            "text": "我负责缓存策略，并记录了命中率。",
        })
        report = self.client.get(f"/api/interview/sessions/{interview_id}/report").json()
        self.assertEqual(report["status"], "finished")
        self.assertEqual(len(report["items"]), 2)
        self.assertEqual(report["items"][0]["skipped"], 1)
        self.assertIsNone(report["items"][0]["answer"])
        self.assertIsNone(report["items"][0]["feedback"])
        self.assertEqual(report["items"][0]["evidence"], [])
        self.assertEqual(report["items"][1]["skipped"], 0)
        self.assertTrue(report["items"][1]["feedback"])
        view = self.client.get(f"/api/interview/sessions/{interview_id}").json()
        self.assertTrue(any(event["kind"] == "skip" for event in view["events"]))
        self.assertIn("已跳过", self.client.get(
            f"/api/interview/sessions/{interview_id}/export.md"
        ).text)

    def test_skipping_last_question_finishes_without_invented_feedback(self):
        interview_id = self.client.post("/api/interview/sessions", json={
            "goal": "后端开发", "max_questions": 1,
        }).json()["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        result = self.client.post(f"/api/interview/sessions/{interview_id}/skip", json={
            "question_number": 1,
        })
        self.assertEqual(result.status_code, 200, result.text)
        report = self.client.get(f"/api/interview/sessions/{interview_id}/report").json()
        self.assertEqual(report["status"], "finished")
        self.assertIn("跳过第 1 题", report["summary"])
        self.assertEqual(report["items"][0]["evidence"], [])
        self.assertIsNone(report["items"][0]["feedback"])
        self.assertFalse(self.client.get("/api/interview/sessions").json()[0]["has_feedback"])

    def test_cancel_interview_keeps_partial_report_and_invalidates_run(self):
        interview_id = self.client.post("/api/interview/sessions", json={"goal": "开发"}).json()["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={"text": "第一题回答"})
        pending = self.client.post(f"/api/interview/sessions/{interview_id}/answers/claim", json={
            "text": "第二题回答", "question_number": 2,
        }).json()
        cancelled = self.client.post(f"/api/interview/sessions/{interview_id}/cancel").json()
        self.assertEqual(cancelled["status"], "cancelled")
        stale = self.client.post(f"/api/interview/sessions/{interview_id}/answers/run", json={
            "version": pending["version"],
        }).json()
        self.assertTrue(stale["interrupted"])
        report = self.client.get(f"/api/interview/sessions/{interview_id}/report").json()
        self.assertEqual(report["items"][0]["answer"], "第一题回答")
        self.assertIsNotNone(report["items"][0]["feedback"])
        self.assertEqual(report["items"][1]["answer"], "第二题回答")
        self.assertIsNone(report["items"][1]["feedback"])
        self.assertIsNone(report["summary"])
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/practice", json={
            "question_number": 1,
        }).status_code, 200)
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/practice", json={
            "question_number": 2,
        }).status_code, 409)
        self.assertEqual(self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
            "text": "继续作答",
        }).status_code, 409)

    def test_session_exports_keep_chat_and_interview_records_separate(self):
        chat_id = self.client.post("/api/chat/sessions").json()["id"]
        self.client.post(f"/api/chat/sessions/{chat_id}/messages", json={"text": "CHAT_PRIVATE"})
        interview_id = self.client.post("/api/interview/sessions", json={
            "goal": "开发岗位", "max_questions": 1,
        }).json()["id"]
        self.client.post(f"/api/interview/sessions/{interview_id}/start")
        self.client.post(f"/api/interview/sessions/{interview_id}/answers", json={
            "text": "INTERVIEW_PRIVATE",
        })
        chat_export = self.client.get(f"/api/chat/sessions/{chat_id}/export.md")
        interview_export = self.client.get(f"/api/interview/sessions/{interview_id}/export.md")
        self.assertEqual(chat_export.status_code, 200)
        self.assertIn("attachment", chat_export.headers["content-disposition"])
        self.assertIn("CHAT_PRIVATE", chat_export.text)
        self.assertNotIn("INTERVIEW_PRIVATE", chat_export.text)
        self.assertIn("INTERVIEW_PRIVATE", interview_export.text)
        self.assertIn("全场总结", interview_export.text)
        self.assertNotIn("CHAT_PRIVATE", interview_export.text)

    def test_deleting_interview_removes_derived_practice_chain(self):
        source_id = self.client.post("/api/interview/sessions", json={
            "goal": "开发岗位", "max_questions": 1,
        }).json()["id"]
        self.client.post(f"/api/interview/sessions/{source_id}/start")
        self.client.post(f"/api/interview/sessions/{source_id}/answers", json={"text": "原回答"})
        practice_id = self.client.post(f"/api/interview/sessions/{source_id}/practice", json={
            "question_number": 1,
        }).json()["id"]
        self.client.post(f"/api/interview/sessions/{practice_id}/start")
        self.client.post(f"/api/interview/sessions/{practice_id}/answers", json={"text": "复练回答"})
        nested_id = self.client.post(f"/api/interview/sessions/{practice_id}/practice", json={
            "question_number": 1,
        }).json()["id"]
        self.assertEqual(self.client.delete(f"/api/interview/sessions/{source_id}").status_code, 200)
        for session_id in (source_id, practice_id, nested_id):
            self.assertEqual(self.client.get(f"/api/interview/sessions/{session_id}").status_code, 404)
        self.assertEqual(self.client.get("/api/interview/sessions").json(), [])

    def test_deleting_chat_leaves_interview_untouched(self):
        chat_id = self.client.post("/api/chat/sessions").json()["id"]
        interview_id = self.client.post("/api/interview/sessions", json={"goal": "开发"}).json()["id"]
        self.assertEqual(self.client.delete(f"/api/chat/sessions/{chat_id}").status_code, 200)
        self.assertEqual(self.client.delete(f"/api/chat/sessions/{chat_id}").status_code, 404)
        self.assertEqual(self.client.get(f"/api/interview/sessions/{interview_id}").status_code, 200)

    def test_material_file_extraction_returns_reviewable_text(self):
        plain = self.client.post("/api/materials/extract", files={
            "file": ("resume.txt", "项目经历：接口优化".encode("utf-8"), "text/plain"),
        })
        self.assertEqual(plain.status_code, 200)
        self.assertEqual(plain.json()["content"], "项目经历：接口优化")

        document = Document()
        document.add_paragraph("项目说明")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "职责"
        table.cell(0, 1).text = "后端开发"
        docx_bytes = BytesIO()
        document.save(docx_bytes)
        word = self.client.post("/api/materials/extract", files={
            "file": ("resume.docx", docx_bytes.getvalue(),
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        })
        self.assertEqual(word.status_code, 200, word.text)
        self.assertIn("项目说明", word.json()["content"])
        self.assertIn("后端开发", word.json()["content"])

        pdf = self.client.post("/api/materials/extract", files={
            "file": ("resume.pdf", self._simple_pdf(), "application/pdf"),
        })
        self.assertEqual(pdf.status_code, 200, pdf.text)
        self.assertIn("Hello PDF", pdf.json()["content"])
        too_long = self.client.post("/api/materials/extract", files={
            "file": ("long.txt", b"x" * 4001, "text/plain"),
        })
        self.assertEqual(too_long.status_code, 422)
        self.assertEqual(self.client.get("/api/materials").json(), [])

    def test_material_file_extraction_accepts_gb18030_text(self):
        for extension in ("txt", "md"):
            with self.subTest(extension=extension):
                response = self.client.post("/api/materials/extract", files={
                    "file": (f"resume.{extension}", "项目经历：接口优化".encode("gb18030"), "text/plain"),
                })
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["content"], "项目经历：接口优化")
        self.assertEqual(self.client.get("/api/materials").json(), [])

    @staticmethod
    def _simple_pdf():
        stream = b"BT /F1 12 Tf 20 100 Td (Hello PDF) Tj ET"
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        ]
        data = b"%PDF-1.4\n"
        offsets = []
        for number, body in enumerate(objects, 1):
            offsets.append(len(data))
            data += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
        xref = len(data)
        data += b"xref\n0 6\n0000000000 65535 f \n"
        data += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets)
        data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF"
        return data
